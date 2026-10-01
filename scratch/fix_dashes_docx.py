"""Replace date/cert separator hyphens with 'to' or commas (no en/em/hyphen separators)."""
import re
from docx import Document

DOCX = "generated_resumes/TarunMannava.docx"
doc = Document(DOCX)

# Date range: "Jan 2025 - May 2026" or "Feb 2022 - Aug 2024"
DATE_RANGE = re.compile(
    r"\b((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4})\s*[-–—]\s*"
    r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}|Present)\b",
    re.I,
)

# Cert issuer separator at end: "Something - Amazon Web Services (AWS)"
CERT_SEP = re.compile(
    r"^(.+?)\s*[-–—]\s*(Amazon Web Services.*|Anthropic.*|Google.*|Microsoft.*)$"
)


def rewrite(text: str) -> str:
    text = DATE_RANGE.sub(r"\1 to \2", text)
    m = CERT_SEP.match(text.strip())
    if m:
        text = f"{m.group(1).strip()}, {m.group(2).strip()}"
    # leftover en/em just in case
    text = text.replace("\u2013", " to ").replace("\u2014", ", ")
    text = text.replace(" -- ", " to ")
    return text


def fix_paragraph(paragraph) -> bool:
    full = "".join(run.text for run in paragraph.runs)
    if not full:
        full = paragraph.text
    new_full = rewrite(full)
    if new_full == full:
        return False
    if paragraph.runs:
        paragraph.runs[0].text = new_full
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_full)
    return True


changed = []
for i, p in enumerate(doc.paragraphs):
    if fix_paragraph(p):
        changed.append((i, p.text))

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if fix_paragraph(p):
                    changed.append(("table", p.text))

doc.save(DOCX)
print(f"Fixed {len(changed)} paragraphs:")
for loc, t in changed:
    print(f"  [{loc}] {t}")
