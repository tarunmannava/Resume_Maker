"""Convert resume LaTeX files in scratch/ to formatted DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

SCRATCH = Path(__file__).resolve().parent

RESUME_GLOB = "Tarun_Mannava_*.tex"
EXTRA_RESUMES = ("resume_django_wildepod.tex", "rewritten_resume.tex")
SKIP_NAMES = {"cover_letter_preamble.tex", "ats_macros.tex"}


def _read_tex(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    macros = SCRATCH / "ats_macros.tex"
    if macros.exists():
        text = text.replace(r"\input{ats_macros.tex}", macros.read_text(encoding="utf-8"))
    return text


def _strip_comments(tex: str) -> str:
    tex = re.sub(r"(?m)^%.*$", "", tex)
    tex = re.sub(r"(?<!\\)%.*$", "", tex, flags=re.MULTILINE)
    return tex


def _extract_body(tex: str) -> str:
    match = re.search(r"\\begin\{document\}(.*)\\end\{document\}", tex, re.DOTALL)
    if not match:
        raise ValueError("No document body found")
    return match.group(1)


def _extract_braced(text: str, start: int) -> tuple[str, int]:
    if start >= len(text) or text[start] != "{":
        raise ValueError(f"Expected '{{' at position {start}")
    depth = 0
    i = start
    while i < len(text):
        ch = text[i]
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1 : i], i + 1
        elif ch == "\\" and i + 1 < len(text):
            i += 1
        i += 1
    raise ValueError("Unbalanced braces")


def _split_macro_args(text: str, macro: str) -> list[str]:
    pattern = re.compile(rf"\\{macro}\s*")
    match = pattern.search(text)
    if not match:
        return []
    pos = match.end()
    args: list[str] = []
    while pos < len(text) and text[pos] in " \t\n\r":
        pos += 1
    while pos < len(text) and text[pos] == "{":
        arg, pos = _extract_braced(text, pos)
        args.append(arg)
        while pos < len(text) and text[pos] in " \t\n\r":
            pos += 1
    return args


def _expand_newcommands(tex: str) -> str:
    pattern = re.compile(
        r"\\newcommand\{\\([a-zA-Z]+)\}\{((?:[^{}]|\{[^{}]*\})*)\}",
        re.DOTALL,
    )
    macros: dict[str, str] = {}
    for match in pattern.finditer(tex):
        macros[match.group(1)] = match.group(2)
    for _ in range(5):
        for name, value in macros.items():
            tex = re.sub(rf"\\{re.escape(name)}\b", lambda _m, v=value: v, tex)
    return tex


def _strip_latex_environments(text: str) -> str:
    text = re.sub(r"\\resumeSubHeadingListStart", "", text)
    text = re.sub(r"\\resumeSubHeadingListEnd", "", text)
    text = re.sub(r"\\resumeItemListStart", "", text)
    text = re.sub(r"\\resumeItemListEnd", "", text)
    text = re.sub(r"\\begin\{itemize\}\[[^\]]*\]", "", text)
    text = re.sub(r"\\begin\{itemize\}", "", text)
    text = re.sub(r"\\end\{itemize\}", "", text)
    return text


def _clean_latex_noise(text: str) -> str:
    text = text.replace(r"\resumeSep", "•")
    text = text.replace(r"\textbullet", "•")
    text = text.replace(r"\textasciitilde", "~")
    text = re.sub(r"\\[1pt]", "\n", text)
    text = re.sub(r"\\[4pt]", "\n", text)
    text = re.sub(r"\\\\", "\n", text)
    text = text.replace(r"\&", "&")
    text = text.replace(r"\%", "%")
    text = text.replace(r"\$", "$")
    text = text.replace(r"\_", "_")
    text = re.sub(r"\\hspace\{[^{}]*\}", " ", text)
    text = re.sub(r"\\vspace\{[^{}]*\}", "", text)
    text = re.sub(r"\\quad\b", "  ", text)
    text = re.sub(r"\$ *\| *\$", " | ", text)
    text = re.sub(r"\\small\b", "", text)
    text = re.sub(r"\\Huge\b", "", text)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r" +", " ", text)
    return text


def _extract_braced_command(text: str, command: str, start: int = 0) -> tuple[str, int] | None:
    match = re.search(rf"\\{command}\b", text[start:])
    if not match:
        return None
    pos = start + match.end()
    while pos < len(text) and text[pos] in " \t\n\r":
        pos += 1
    if pos >= len(text) or text[pos] != "{":
        return None
    content, end = _extract_braced(text, pos)
    return content, end


def _add_formatted_runs(paragraph: Paragraph, text: str, *, base_size: float = 10.5) -> None:
    if not text.strip():
        return

    pos = 0
    while pos < len(text):
        if pos >= len(text):
            break

        candidates: list[tuple[int, str, str | None, str | None]] = []
        for cmd, kind in (
            ("href", "href"),
            ("textbf", "bold"),
            ("textit", "italic"),
            ("emph", "italic"),
        ):
            found = _extract_braced_command(text, cmd, pos)
            if found and (start := text.find(f"\\{cmd}", pos)) != -1:
                if cmd == "href":
                    pos_href = start + len(cmd) + 1
                    while pos_href < len(text) and text[pos_href] in " \t\n\r":
                        pos_href += 1
                    if pos_href < len(text) and text[pos_href] == "{":
                        url, after_url = _extract_braced(text, pos_href)
                        while after_url < len(text) and text[after_url] in " \t\n\r":
                            after_url += 1
                        if after_url < len(text) and text[after_url] == "{":
                            label, _ = _extract_braced(text, after_url)
                            candidates.append((start, "href", url, label))
                else:
                    content, _ = found
                    candidates.append((start, kind, content, None))

        next_cmd = re.search(r"\\(small|Huge)\b", text[pos:])
        if next_cmd:
            candidates.append((pos + next_cmd.start(), "noop", None, None))

        if not candidates:
            chunk_end = len(text)
            next_special = re.search(r"\\(href|textbf|textit|emph|small|Huge)\b", text[pos:])
            if next_special:
                chunk_end = pos + next_special.start()
            chunk = text[pos:chunk_end]
            if chunk:
                cleaned = _clean_latex_noise(chunk)
                if cleaned:
                    run = paragraph.add_run(cleaned)
                    run.font.size = Pt(base_size)
            pos = chunk_end
            continue

        candidates.sort(key=lambda item: item[0])
        start, kind, a, b = candidates[0]

        if start > pos:
            chunk = text[pos:start]
            cleaned = _clean_latex_noise(chunk)
            if cleaned:
                run = paragraph.add_run(cleaned)
                run.font.size = Pt(base_size)

        if kind == "href" and a and b:
            _add_hyperlink(paragraph, a, _clean_latex_noise(b), size=base_size)
            pos = start
            m = re.search(r"\\href\b", text[pos:])
            if m is None:
                pos = start + 1
                continue
            pos += m.end()
            while pos < len(text) and text[pos] in " \t\n\r":
                pos += 1
            if pos < len(text) and text[pos] == "{":
                _, pos = _extract_braced(text, pos)
            while pos < len(text) and text[pos] in " \t\n\r":
                pos += 1
            if pos < len(text) and text[pos] == "{":
                _, pos = _extract_braced(text, pos)
        elif kind in {"bold", "italic"} and a:
            run = paragraph.add_run(_clean_latex_noise(a).strip())
            run.bold = kind == "bold"
            run.italic = kind == "italic"
            run.font.size = Pt(base_size)
            pos = start
            cmd_name = "textbf" if kind == "bold" else "textit"
            m = re.search(rf"\\{cmd_name}\b", text[pos:])
            if m is None and kind == "italic":
                m = re.search(r"\\emph\b", text[pos:])
            if m is None:
                pos = start + 1
                continue
            pos += m.end()
            while pos < len(text) and text[pos] in " \t\n\r":
                pos += 1
            if pos < len(text) and text[pos] == "{":
                _, pos = _extract_braced(text, pos)
        elif kind == "noop":
            m = re.search(r"\\(small|Huge)\b", text[start:])
            pos = start + m.end() if m else start + 1
        else:
            pos = start + 1


def _add_hyperlink(paragraph: Paragraph, url: str, label: str, *, size: float = 10.5) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = label
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_section_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.small_caps = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    # underline rule effect via bottom border
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    _add_formatted_runs(p, text)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.15)


def _parse_heading(doc: Document, body: str) -> None:
    center = re.search(r"\\begin\{center\}(.*?)\\end\{center\}", body, re.DOTALL)
    if not center:
        return
    block = center.group(1)
    lines = [ln.strip() for ln in re.split(r"\\\\(?:\[[^\]]*\])?", block) if ln.strip()]
    for i, line in enumerate(lines):
        if i == 0:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_formatted_runs(p, line, base_size=20 if "Huge" in line else 14)
            for run in p.runs:
                if "Tarun" in run.text:
                    run.bold = True
                    run.font.size = Pt(20)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_formatted_runs(p, line, base_size=10.5)
            if line.startswith(r"\textit") or "engineer" in line.lower():
                for run in p.runs:
                    run.italic = True


def _parse_subheading(doc: Document, args: list[str]) -> None:
    if len(args) < 4:
        return
    org, title, start, end = args[:4]
    p = doc.add_paragraph()
    r1 = p.add_run(_clean_latex_noise(org).strip())
    r1.bold = True
    r1.font.size = Pt(10.5)
    p.add_run("\n")
    r2 = p.add_run(_clean_latex_noise(title).strip())
    r2.italic = True
    r2.font.size = Pt(10.5)
    p.add_run("\n")
    r3 = p.add_run(f"{_clean_latex_noise(start).strip()} – {_clean_latex_noise(end).strip()}")
    r3.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)


def _parse_project(doc: Document, args: list[str]) -> None:
    if len(args) == 3:
        name, tech, dates = args
        p = doc.add_paragraph()
        r1 = p.add_run(_clean_latex_noise(name))
        r1.bold = True
        r1.font.size = Pt(10.5)
        p.add_run("\n")
        r2 = p.add_run(_clean_latex_noise(tech))
        r2.italic = True
        r2.font.size = Pt(10.5)
        p.add_run("\n")
        r3 = p.add_run(_clean_latex_noise(dates))
        r3.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(2)
    elif len(args) == 2:
        name, tech = args
        p = doc.add_paragraph()
        r1 = p.add_run(_clean_latex_noise(name))
        r1.bold = True
        r1.font.size = Pt(10.5)
        p.add_run("\n")
        r2 = p.add_run(_clean_latex_noise(tech))
        r2.italic = True
        r2.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(2)
    elif len(args) == 1:
        p = doc.add_paragraph()
        r = p.add_run(_clean_latex_noise(args[0]))
        r.bold = True
        r.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(2)


def _parse_item_block(doc: Document, block: str) -> None:
    block = _strip_latex_environments(block)
    for item_match in re.finditer(r"\\item\b", block):
        pos = item_match.end()
        while pos < len(block) and block[pos] in " \t\n\r":
            pos += 1
        if pos >= len(block):
            continue
        if block[pos] == "{":
            try:
                content, _ = _extract_braced(block, pos)
            except ValueError:
                continue
        else:
            line_end = block.find("\n", pos)
            content = block[pos:line_end if line_end != -1 else len(block)]
        for line in re.split(r"\\\\(?:\[[^\]]*\])?", content):
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph()
            _add_formatted_runs(p, re.sub(r"\\small\b", "", line))
            p.paragraph_format.space_after = Pt(1)


def _parse_section_content(doc: Document, content: str, *, section_title: str) -> None:
    title = section_title.upper()
    if title in {"SKILLS", "CERTIFICATIONS"}:
        _parse_item_block(doc, content)
        return

    content = _strip_latex_environments(content)
    macros = [
        "resumeSubheading",
        "resumeProject",
        "resumeProjectHeading",
        "resumeItem",
    ]
    pos = 0
    while pos < len(content):
        next_macro = None
        next_pos = len(content)
        for macro in macros:
            m = re.search(rf"\\{macro}\b", content[pos:])
            if m:
                abs_pos = pos + m.start()
                if abs_pos < next_pos:
                    next_pos = abs_pos
                    next_macro = macro
        if next_macro is None:
            break
        args = _split_macro_args(content[next_pos:], next_macro)
        if next_macro == "resumeSubheading":
            _parse_subheading(doc, args)
        elif next_macro in {"resumeProject", "resumeProjectHeading"}:
            _parse_project(doc, args)
        elif next_macro == "resumeItem" and args:
            _add_bullet(doc, args[0])
        m = re.search(rf"\\{next_macro}\b", content[next_pos:])
        end = next_pos + m.end() if m else next_pos
        pos = end
        for _ in args:
            if pos < len(content) and content[pos] == "{":
                _, pos = _extract_braced(content, pos)
            while pos < len(content) and content[pos] in " \t\n\r":
                pos += 1


def convert_tex_to_docx(tex_path: Path, out_path: Path) -> None:
    tex = _expand_newcommands(_strip_comments(_read_tex(tex_path)))
    body = _extract_body(tex)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    _parse_heading(doc, body)

    for section_match in re.finditer(r"\\section\{([^{}]+)\}(.*?)(?=\\section\{|\Z)", body, re.DOTALL):
        title = section_match.group(1).strip()
        content = section_match.group(2)
        _add_section_heading(doc, title)
        _parse_section_content(doc, content, section_title=title)

    doc.save(out_path)
    print(f"Wrote {out_path.name}")


def main() -> None:
    paths: list[Path] = []
    for pattern in [RESUME_GLOB, *EXTRA_RESUMES]:
        paths.extend(SCRATCH.glob(pattern) if "*" in pattern else [SCRATCH / pattern])

    seen: set[Path] = set()
    for tex_path in sorted(paths):
        if tex_path.name in SKIP_NAMES or "Cover_Letter" in tex_path.name:
            continue
        if tex_path in seen:
            continue
        seen.add(tex_path)
        out_path = tex_path.with_suffix(".docx")
        convert_tex_to_docx(tex_path, out_path)


if __name__ == "__main__":
    main()
