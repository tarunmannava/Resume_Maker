"""Put PROJECTS section above EXPERIENCE in TarunMannava.docx."""
from docx import Document

path = "generated_resumes/TarunMannava.docx"
doc = Document(path)
body = doc.element.body
children = list(body)  # includes sectPr at end

# Current mess after failed moves:
# 0-11: header..skills
# 12-37: project content REVERSED
# 38: PROJECTS
# 39-61: EXPERIENCE..certs
# 62: sectPr

skills_end = 11
proj_content_rev = children[12:38]
projects_header = children[38]
rest = children[39:]  # EXPERIENCE through sectPr

projects_correct = [projects_header] + list(reversed(proj_content_rev))

new_order = children[: skills_end + 1] + projects_correct + rest

# Rebuild body
for child in list(body):
    body.remove(child)
for child in new_order:
    body.append(child)

doc.save(path)

# Verify
doc2 = Document(path)
print("Section headers + next line:")
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    if t in (
        "SUMMARY",
        "SKILLS",
        "PROJECTS",
        "EXPERIENCE",
        "EDUCATION",
        "CERTIFICATIONS",
    ):
        nxt = doc2.paragraphs[i + 1].text[:100] if i + 1 < len(doc2.paragraphs) else ""
        print(f"{i} {t}")
        print(f"   -> {nxt}")
