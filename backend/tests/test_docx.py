from pathlib import Path

from docx import Document

from backend.app.services.docx import (
    _clean_latex_noise,
    _expand_newcommands,
    _has_section_content,
    convert_tex_to_docx,
)

RESUME_TEX = r"""
\documentclass{article}
\newcommand{\resumeItem}[1]{\item #1}
\newcommand{\resumeSubheading}[4]{\item \textbf{#1} \\ \textit{#2} \\ #3 -- #4}
\newcommand{\resumeProject}[2]{\item[] \textbf{#1} \\ \textit{#2}}
\newcommand{\resumeSubHeadingListStart}{\begin{itemize}}
\newcommand{\resumeSubHeadingListEnd}{\end{itemize}}
\newcommand{\resumeItemListStart}{\begin{itemize}}
\newcommand{\resumeItemListEnd}{\end{itemize}}
\begin{document}
\begin{center}
{\Huge \textbf{Test Candidate}} \\
test@example.com
\end{center}
\section{EXPERIENCE}
\resumeSubHeadingListStart
\resumeSubheading{Example Corp}{Software Engineer}{Jan 2020}{Present}
\resumeItemListStart
\resumeItem{Built things.}
\resumeItemListEnd
\resumeSubHeadingListEnd
\section{PROJECTS}
\resumeSubHeadingListStart
\resumeProject{Sample Project}{Python, FastAPI}
\resumeItemListStart
\resumeItem{Did project work.}
\resumeItemListEnd
\resumeSubHeadingListEnd
\end{document}
"""


def test_docx_project_has_no_date_range(tmp_path: Path):
    """Projects intentionally have no date range (unlike Experience/Education):
    \\resumeProject{title}{tech stack} only takes 2 arguments, and the DOCX
    output must render just the title + tech stack with no stray date text."""
    tex_path = tmp_path / "resume.tex"
    tex_path.write_text(RESUME_TEX, encoding="utf-8")
    out_path = tmp_path / "resume.docx"

    convert_tex_to_docx(tex_path, out_path)

    doc = Document(str(out_path))
    project_line = next(p.text for p in doc.paragraphs if "Sample Project" in p.text)
    assert project_line == "Sample Project | Python, FastAPI"

    experience_line = next(p.text for p in doc.paragraphs if "Example Corp" in p.text)
    assert "Jan 2020" in experience_line
    assert "Present" in experience_line


def test_docx_project_heading_with_href_and_underline(tmp_path: Path):
    tex_content = r"""
\documentclass{article}
\newcommand{\resumeItem}[1]{\item #1}
\newcommand{\resumeProjectHeading}[2]{\item[] #1 \\ \textit{#2}}
\newcommand{\resumeSubHeadingListStart}{\begin{itemize}}
\newcommand{\resumeSubHeadingListEnd}{\end{itemize}}
\newcommand{\resumeItemListStart}{\begin{itemize}}
\newcommand{\resumeItemListEnd}{\end{itemize}}
\begin{document}
\section{PROJECTS}
\resumeSubHeadingListStart
\resumeProjectHeading{\textbf{BakedBeacon -- Skills Platform} $|$ \href{https://skillbeacon-six.vercel.app/}{\underline{Link}}}{Python, FastAPI, React}
\resumeItemListStart
\resumeItem{Engineered backend microservices.}
\resumeItemListEnd
\resumeSubHeadingListEnd
\section{CERTIFICATIONS}
\resumeSubHeadingListStart
\resumeSubHeadingListEnd
\end{document}
"""
    tex_path = tmp_path / "resume_links.tex"
    tex_path.write_text(tex_content, encoding="utf-8")
    out_path = tmp_path / "resume_links.docx"

    convert_tex_to_docx(tex_path, out_path)

    doc = Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # 1. Ensure raw \href or \underline didn't leak into the plain text
    assert r"\href" not in full_text
    assert r"\underline" not in full_text
    assert "BakedBeacon -- Skills Platform" in full_text
    assert "Link" in full_text

    # 2. Ensure empty CERTIFICATIONS section was omitted
    assert "CERTIFICATIONS" not in full_text


def test_docx_header_math_mode_separators_and_comments(tmp_path: Path):
    tex_content = r"""
\documentclass{article}
\begin{document}
\begin{center}
{\Huge \textbf{Tarun Mannava}} \\
Seattle, WA ~$\bullet$~ +1 (656) 203-7074 ~$\bullet$~ \href{mailto:mannava.t6@gmail.com}{mannava.t6@gmail.com} ~$\bullet$~ \href{https://linkedin.com/in/tarunmannava}{LinkedIn}%
\end{center}
\section{SUMMARY}
Software Engineer with 3+ years experience.
\end{document}
"""
    tex_path = tmp_path / "resume_header.tex"
    tex_path.write_text(tex_content, encoding="utf-8")
    out_path = tmp_path / "resume_header.docx"

    convert_tex_to_docx(tex_path, out_path)

    doc = Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "$$" not in full_text
    assert "$" not in full_text
    assert "%" not in full_text
    assert "Seattle, WA • +1 (656) 203-7074 • mannava.t6@gmail.com • LinkedIn" in full_text


def test_docx_header_hspace_textbullet_links(tmp_path: Path):
    tex_content = r"""
\documentclass{article}
\begin{document}
\begin{center}
{\Huge \textbf{Tarun Mannava}} \\
Seattle, WA \hspace{0.4em}\textbullet\hspace{0.4em}+1 (656) 203-7074\hspace{0.4em}\textbullet\hspace{0.4em}\href{mailto:mannava.tarun34@gmail.com}{mannava.tarun34@gmail.com}\hspace{0.4em}\textbullet\hspace{0.4em}\href{https://linkedin.com/in/tarunmannava}{LinkedIn}\hspace{0.4em}\textbullet\hspace{0.4em}\href{https://github.com/tarunmannava}{GitHub}
\end{center}
\section{SUMMARY}
Software Engineer.
\end{document}
"""
    tex_path = tmp_path / "resume_hspace.tex"
    tex_path.write_text(tex_content, encoding="utf-8")
    out_path = tmp_path / "resume_hspace.docx"

    convert_tex_to_docx(tex_path, out_path)

    doc = Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "hspace" not in full_text
    assert "0.4em" not in full_text
    assert "textbullet" not in full_text
    assert "href" not in full_text
    assert "Seattle, WA • +1 (656) 203-7074 • mannava.tarun34@gmail.com • LinkedIn • GitHub" in full_text


def test_docx_education_and_skills_rendering(tmp_path: Path):
    tex_content = r"""
\documentclass{article}
\newcommand{\resumeSubheading}[4]{\item \textbf{#1} \\ \textit{#2} \\ #3 -- #4}
\begin{document}
\section{SKILLS}
\begin{itemize}
\item{
  \textbf{Languages:} Python, TypeScript, Java \\[1pt]
  \textbf{Backend:} FastAPI, Spring Boot
}
\end{itemize}
\section{EDUCATION}
\resumeSubheading
  {University of South Florida}{Master of Science, Computer Science}{Aug 2024}{May 2026}
\end{document}
"""
    tex_path = tmp_path / "resume_edu.tex"
    tex_path.write_text(tex_content, encoding="utf-8")
    out_path = tmp_path / "resume_edu.docx"

    convert_tex_to_docx(tex_path, out_path)

    doc = Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "SKILLS" in full_text
    assert "Languages: Python, TypeScript, Java" in full_text
    assert "Backend: FastAPI, Spring Boot" in full_text
    assert "EDUCATION" in full_text
    assert "University of South Florida" in full_text
    assert "Master of Science, Computer Science" in full_text
    assert "Aug 2024 – May 2026" in full_text


def test_docx_real_resume_file_conversion(tmp_path: Path):
    source_resume = Path("Tarun_Mannava_Resume.tex")
    if not source_resume.exists():
        return
    out_path = tmp_path / "Tarun_Mannava_Resume.docx"
    convert_tex_to_docx(source_resume, out_path)
    assert out_path.exists()
    assert out_path.stat().st_size > 0

    doc = Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Tarun Mannava" in full_text
    assert "Cognizant Technology Solutions" in full_text
    assert "SkillBeacon" in full_text
    assert "CERTIFICATIONS" in full_text
    assert "AWS Certified Cloud Practitioner" in full_text
