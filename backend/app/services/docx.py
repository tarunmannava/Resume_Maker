"""Convert resume LaTeX files to beautifully formatted ATS-compliant DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


def _read_tex(path: Path) -> str:
    """Read a LaTeX file and expand any \\input{...} directives found relative to the file."""
    text = path.read_text(encoding="utf-8")
    
    def replace_input(m: re.Match) -> str:
        input_target = m.group(1).strip()
        if not input_target.endswith(".tex"):
            input_target += ".tex"
        input_path = path.parent / input_target
        if input_path.exists():
            return input_path.read_text(encoding="utf-8")
        return ""
    
    text = re.sub(r"\\input\{([^}]+)\}", replace_input, text)
    return text


def _strip_comments(tex: str) -> str:
    r"""Strip LaTeX comments (%...) while preserving escaped \%."""
    tex = re.sub(r"(?m)^%.*$", "", tex)
    tex = re.sub(r"(?<!\\)%.*$", "", tex, flags=re.MULTILINE)
    return tex


def _extract_body(tex: str) -> str:
    """Extract content within \\begin{document}...\\end{document} or return full text if missing."""
    match = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", tex, re.DOTALL)
    if match:
        return match.group(1)
    return tex


def _extract_braced(text: str, start: int) -> tuple[str, int]:
    """Extract a balanced {...} group starting at `start`. Returns (content, next_index)."""
    if start >= len(text) or text[start] != "{":
        raise ValueError(f"Expected '{{' at position {start}")
    depth = 0
    i = start
    while i < len(text):
        ch = text[i]
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1 : i], i + 1
        elif ch == "\\" and i + 1 < len(text):
            i += 1
        i += 1
    # Fallback if unbalanced
    return text[start + 1 :], len(text)


def _expand_newcommands(tex: str) -> str:
    """Expand simple 0-argument \\newcommand macros defined in preamble or text."""
    pattern = re.compile(
        r"\\newcommand\{\\([a-zA-Z]+)\}\{((?:[^{}]|\{[^{}]*\})*)\}",
        re.DOTALL,
    )
    macros: dict[str, str] = {}
    for match in pattern.finditer(tex):
        macros[match.group(1)] = match.group(2)
    
    # Expand iteratively up to 5 passes for nested macros
    for _ in range(5):
        for name, value in macros.items():
            tex = re.sub(rf"\\{re.escape(name)}\b", lambda _m, v=value: v, tex)
    return tex


def _strip_latex_environments(text: str) -> str:
    """Strip common list environment wrappers."""
    text = re.sub(r"\\resumeSubHeadingListStart", "", text)
    text = re.sub(r"\\resumeSubHeadingListEnd", "", text)
    text = re.sub(r"\\resumeItemListStart", "", text)
    text = re.sub(r"\\resumeItemListEnd", "", text)
    text = re.sub(r"\\begin\{itemize\}(?:\[[^\]]*\])?", "", text)
    text = re.sub(r"\\end\{itemize\}", "", text)
    return text


def _clean_latex_noise(text: str) -> str:
    """Clean LaTeX macros, math mode symbols, and formatting noise to return human text."""
    # 1. Math mode symbols & bullets
    text = re.sub(r"\$\s*\\(?:bullet|cdot|circ|diamond|star)\s*\$", " • ", text)
    text = re.sub(r"\$\s*\\(?:vert|mid|\|)\s*\$", " | ", text)
    text = re.sub(r"\$\s*\\sim\s*\$", " ~ ", text)
    text = re.sub(r"\$\s*\|\s*\$", " | ", text)
    text = re.sub(r"\$\s*\$", " ", text)
    
    # 2. Text bullets & symbol macros
    text = text.replace(r"\resumeSep", "•")
    text = text.replace(r"\textbullet", "•")
    text = text.replace(r"\bullet", "•")
    text = text.replace(r"\cdot", "•")
    text = text.replace(r"\vert", "|")
    text = text.replace(r"\mid", "|")
    text = text.replace(r"\textbar", "|")
    text = text.replace(r"\textasciitilde", "~")
    text = text.replace(r"\sim", "~")
    
    # 3. LaTeX newlines (\\ or \\[4pt] or \\[0.4em])
    text = re.sub(r"\\\\(?:\[[^\]]*\])?", "\n", text)
    
    # 4. Spacing and structural commands
    text = re.sub(r"\\(?:hspace|vspace)\*?\{[^}]*\}", " ", text)
    text = re.sub(r"\\(?:fontsize|selectfont|raggedright|raggedbottom|pagestyle|geometry|setstretch)\*?(?:\{[^}]*\})*", "", text)
    text = re.sub(r"\\quad\b", "  ", text)
    text = re.sub(r"\\qquad\b", "    ", text)
    text = text.replace("~", " ")
    
    # 5. Escaped characters
    text = text.replace(r"\&", "&")
    text = text.replace(r"\_", "_")
    text = text.replace(r"\#", "#")
    text = text.replace(r"\$", "$")
    
    # 6. Font commands
    text = re.sub(r"\\(small|Huge|huge|large|Large|normalsize|scshape)\b", "", text)
    
    # 7. Unwrapping formatting macros
    for _ in range(4):
        text = re.sub(r"\\(?:underline|textbf|textit|emph|textsc|texttt)\{([^{}]*)\}", r"\1", text)
        text = re.sub(r"\\href\{[^{}]*\}\{([^{}]*)\}", r"\1", text)
        text = re.sub(r"\\url\{([^{}]*)\}", r"\1", text)
        
    # 8. Strip leftover macros and braces
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?(?:\{[^{}]*\})*", "", text)
    text = re.sub(r"\\[a-zA-Z]+\*?", "", text)
    text = re.sub(r"\$+", " ", text)
    text = re.sub(r"(?m)\s*\\?%\s*$", "", text)
    text = re.sub(r"(?m)^\s*\\?%\s*", "", text)
    text = text.replace(r"\%", "%")
    text = re.sub(r"[{}]", "", text)
    
    # 9. Clean excessive spaces around delimiters
    text = re.sub(r"[ \t]*•[ \t]*", " • ", text)
    text = re.sub(r"[ \t]*\|[ \t]*", " | ", text)
    text = re.sub(r"[ \t]+", " ", text)
    
    has_leading = text.startswith(" ")
    has_trailing = text.endswith(" ")
    text = text.strip()
    if not text:
        return ""
    if has_leading:
        text = " " + text
    if has_trailing:
        text = text + " "
    return text


def _extract_braced_command(text: str, command: str, start: int = 0) -> tuple[str, int] | None:
    """Find \\command{...} starting at or after `start` and extract its braced argument."""
    match = re.search(rf"\\{command}\b", text[start:])
    if not match:
        return None
    pos = start + match.end()
    while pos < len(text) and text[pos] in " \t\n\r":
        pos += 1
    if pos >= len(text) or text[pos] != "{":
        return None
    content, end = _extract_braced(text, pos)
    return content, end


def _add_hyperlink(paragraph: Paragraph, url: str, label: str, *, size: float = 9.5) -> None:
    """Add a clickable external hyperlink with blue font and underline styling."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    
    # Underline
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    
    # Color #0563C1 (standard Microsoft Word hyperlink blue)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    
    # Font Size (in half-points)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = label
    # Preserve whitespace in xml:space
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_formatted_runs(
    paragraph: Paragraph,
    text: str,
    *,
    base_size: float = 9.5,
    default_bold: bool = False,
    default_italic: bool = False,
) -> None:
    """Recursively parse and append formatted runs to a Word paragraph."""
    # Pre-clean known non-visual commands
    text = re.sub(r"\\(small|Huge|huge|large|Large|normalsize|scshape)\b", "", text)
    text = text.strip()
    if not text:
        return

    pos = 0
    while pos < len(text):
        if pos >= len(text):
            break

        candidates: list[tuple[int, str, str | None, str | None]] = []
        for cmd, kind in (
            ("href", "href"),
            ("url", "url"),
            ("textbf", "bold"),
            ("textit", "italic"),
            ("emph", "italic"),
            ("underline", "underline"),
        ):
            found = _extract_braced_command(text, cmd, pos)
            if found and (start := text.find(f"\\{cmd}", pos)) != -1:
                if cmd == "href":
                    pos_href = start + len(cmd) + 1
                    while pos_href < len(text) and text[pos_href] in " \t\n\r":
                        pos_href += 1
                    if pos_href < len(text) and text[pos_href] == "{":
                        url, after_url = _extract_braced(text, pos_href)
                        while after_url < len(text) and text[after_url] in " \t\n\r":
                            after_url += 1
                        if after_url < len(text) and text[after_url] == "{":
                            label, _ = _extract_braced(text, after_url)
                            candidates.append((start, "href", url, label))
                elif cmd == "url":
                    url_content, _ = found
                    candidates.append((start, "url", url_content, url_content))
                else:
                    content, _ = found
                    candidates.append((start, kind, content, None))

        # Check for structural no-ops
        next_cmd = re.search(r"\\(small|Huge|huge|large|Large|normalsize|scshape)\b", text[pos:])
        if next_cmd:
            candidates.append((pos + next_cmd.start(), "noop", None, None))

        if not candidates:
            chunk = text[pos:]
            if chunk:
                cleaned = _clean_latex_noise(chunk)
                if cleaned:
                    run = paragraph.add_run(cleaned)
                    run.font.size = Pt(base_size)
                    if default_bold:
                        run.bold = True
                    if default_italic:
                        run.italic = True
            break

        candidates.sort(key=lambda item: item[0])
        start, kind, a, b = candidates[0]

        # Add any plain text preceding the command
        if start > pos:
            chunk = text[pos:start]
            cleaned = _clean_latex_noise(chunk)
            if cleaned:
                run = paragraph.add_run(cleaned)
                run.font.size = Pt(base_size)
                if default_bold:
                    run.bold = True
                if default_italic:
                    run.italic = True

        if kind in {"href", "url"} and a:
            url_target = a.strip()
            label_text = _clean_latex_noise(b if b else a).strip()
            if not label_text:
                label_text = url_target
            _add_hyperlink(paragraph, url_target, label_text, size=base_size)
            
            # Advance pos past the href/url command
            pos = start
            cmd_name = "href" if kind == "href" else "url"
            m = re.search(rf"\\{cmd_name}\b", text[pos:])
            if m is None:
                pos = start + 1
                continue
            pos += m.end()
            while pos < len(text) and text[pos] in " \t\n\r":
                pos += 1
            if pos < len(text) and text[pos] == "{":
                _, pos = _extract_braced(text, pos)
            if kind == "href":
                while pos < len(text) and text[pos] in " \t\n\r":
                    pos += 1
                if pos < len(text) and text[pos] == "{":
                    _, pos = _extract_braced(text, pos)
                    
        elif kind in {"bold", "italic", "underline"} and a:
            inner_text = a
            # If inner text contains nested commands (e.g. href or bold inside italic), recurse
            if any(c in inner_text for c in (r"\href", r"\url", r"\textbf", r"\textit", r"\emph", r"\underline")):
                _add_formatted_runs(
                    paragraph,
                    inner_text,
                    base_size=base_size,
                    default_bold=(kind == "bold" or default_bold),
                    default_italic=(kind == "italic" or default_italic),
                )
            else:
                cleaned = _clean_latex_noise(inner_text)
                if cleaned:
                    run = paragraph.add_run(cleaned)
                    run.font.size = Pt(base_size)
                    run.bold = kind == "bold" or default_bold
                    run.italic = kind == "italic" or default_italic
                    run.underline = kind == "underline"
            
            # Advance pos past this command
            pos = start
            cmd_name = "textbf" if kind == "bold" else ("textit" if kind == "italic" else "underline")
            m = re.search(rf"\\{cmd_name}\b", text[pos:])
            if m is None and kind == "italic":
                m = re.search(r"\\emph\b", text[pos:])
            if m is None:
                pos = start + 1
                continue
            pos += m.end()
            while pos < len(text) and text[pos] in " \t\n\r":
                pos += 1
            if pos < len(text) and text[pos] == "{":
                _, pos = _extract_braced(text, pos)
                
        elif kind == "noop":
            m = re.search(r"\\(small|Huge|huge|large|Large|normalsize|scshape)\b", text[start:])
            pos = start + m.end() if m else start + 1
        else:
            pos = start + 1


def _add_section_heading(doc: Document, title: str) -> None:
    """Render an uppercase section heading with a crisp single bottom border."""
    p = doc.add_paragraph()
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(9.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_bullet(doc: Document, text: str) -> None:
    """Add a compact bullet point item with standard ATS indent and spacing."""
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    _add_formatted_runs(p, text, base_size=9.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.left_indent = Inches(0.15)


def _parse_heading(doc: Document, body: str) -> None:
    """Parse candidate header (name, contact line, optional professional summary)."""
    center = re.search(r"\\begin\{center\}(.*?)\\end\{center\}", body, re.DOTALL)
    if not center:
        return
    block = center.group(1)
    lines = [ln.strip() for ln in re.split(r"\\\\(?:\[[^\]]*\])?", block) if ln.strip()]
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4) if i == len(lines) - 1 else Pt(0)

        if i == 0:
            # Candidate Name line
            _add_formatted_runs(p, line, base_size=18.0, default_bold=True)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(18.0)
        else:
            # Contact or Summary line
            _add_formatted_runs(p, line, base_size=9.5)


def _split_macro_args(text: str, macro: str) -> list[str]:
    """Extract positional {...} arguments for a macro call."""
    pattern = re.compile(rf"\\{macro}\s*")
    match = pattern.search(text)
    if not match:
        return []
    pos = match.end()
    args: list[str] = []
    while pos < len(text) and text[pos] in " \t\n\r":
        pos += 1
    while pos < len(text) and text[pos] == "{":
        arg, pos = _extract_braced(text, pos)
        args.append(arg)
        while pos < len(text) and text[pos] in " \t\n\r":
            pos += 1
    return args


def _parse_subheading(doc: Document, args: list[str], *, section_title: str = "") -> None:
    """Parse experience/education subheading with right-aligned dates."""
    if len(args) < 4:
        return
    org, title, start, end = args[:4]
    
    if section_title.upper() == "EDUCATION":
        # Line 1: Institution (Left) + Dates (Right)
        p1 = doc.add_paragraph()
        p1.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
        p1.paragraph_format.space_before = Pt(1)
        _add_formatted_runs(p1, org, base_size=9.5, default_bold=True)
        p1.add_run("\t")
        date_str = f"{_clean_latex_noise(start).strip()} – {_clean_latex_noise(end).strip()}"
        r_date = p1.add_run(date_str)
        r_date.bold = True
        r_date.font.size = Pt(9.5)
        
        # Line 2: Degree (Italic)
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        _add_formatted_runs(p2, title, base_size=9.5, default_italic=True)
    else:
        # Experience: Company | Title (Left) + Dates (Right)
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(1)
        
        _add_formatted_runs(p, org, base_size=9.5, default_bold=True)
        if title.strip():
            p.add_run(" | ")
            _add_formatted_runs(p, title, base_size=9.5, default_bold=True, default_italic=True)
            
        p.add_run("\t")
        date_str = f"{_clean_latex_noise(start).strip()} – {_clean_latex_noise(end).strip()}"
        r_date = p.add_run(date_str)
        r_date.bold = True
        r_date.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(1)


def _parse_project(doc: Document, args: list[str]) -> None:
    """Parse project header line: Title | [Link] | Tech Stack."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)

    if len(args) >= 1:
        name = args[0]
        _add_formatted_runs(p, name, base_size=9.5, default_bold=True)

    if len(args) >= 2:
        tech = args[1].strip()
        if tech:
            p.add_run(" | ")
            _add_formatted_runs(p, tech, base_size=9.5, default_italic=True)

    p.paragraph_format.space_after = Pt(1)


def _parse_item_block(doc: Document, block: str) -> None:
    """Parse list items in Skills, Certifications, or Awards sections."""
    block = _strip_latex_environments(block)
    
    # Check if block uses \item commands
    items = list(re.finditer(r"\\item\b(?:\[[^\]]*\])?", block))
    if items:
        for idx, item_match in enumerate(items):
            pos = item_match.end()
            while pos < len(block) and block[pos] in " \t\n\r":
                pos += 1
            if pos >= len(block):
                continue
            if block[pos] == "{":
                try:
                    content, _ = _extract_braced(block, pos)
                except ValueError:
                    content = block[pos:]
            else:
                next_pos = items[idx + 1].start() if idx + 1 < len(items) else len(block)
                content = block[pos:next_pos]
                
            for line in re.split(r"\\\\(?:\[[^\]]*\])?|\n", content):
                line = line.strip()
                if not line:
                    continue
                p = doc.add_paragraph()
                _add_formatted_runs(p, line, base_size=9.5)
                p.paragraph_format.space_after = Pt(1)
    else:
        # Plain lines
        for line in re.split(r"\\\\(?:\[[^\]]*\])?|\n", block):
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph()
            _add_formatted_runs(p, line, base_size=9.5)
            p.paragraph_format.space_after = Pt(1)


def _parse_section_content(doc: Document, content: str, *, section_title: str) -> None:
    """Parse inner content of a resume section."""
    title = section_title.upper()
    if any(k in title for k in ("SKILLS", "CERTIFICATIONS", "AWARDS", "HONORS", "PUBLICATIONS")):
        _parse_item_block(doc, content)
        return

    content = _strip_latex_environments(content)
    macros = [
        "resumeSubheading",
        "resumeProject",
        "resumeProjectHeading",
        "resumeItem",
    ]
    pos = 0
    while pos < len(content):
        next_macro = None
        next_pos = len(content)
        for macro in macros:
            m = re.search(rf"\\{macro}\b", content[pos:])
            if m:
                abs_pos = pos + m.start()
                if abs_pos < next_pos:
                    next_pos = abs_pos
                    next_macro = macro
        if next_macro is None:
            break
        args = _split_macro_args(content[next_pos:], next_macro)
        if next_macro == "resumeSubheading":
            _parse_subheading(doc, args, section_title=section_title)
        elif next_macro in {"resumeProject", "resumeProjectHeading"}:
            _parse_project(doc, args)
        elif next_macro == "resumeItem" and args:
            _add_bullet(doc, args[0])
        m = re.search(rf"\\{next_macro}\b", content[next_pos:])
        end = next_pos + m.end() if m else next_pos
        pos = end
        for _ in args:
            if pos < len(content) and content[pos] == "{":
                _, pos = _extract_braced(content, pos)
            while pos < len(content) and content[pos] in " \t\n\r":
                pos += 1


def _has_section_content(content: str) -> bool:
    """Determine if a section contains actual text rather than empty LaTeX boilerplate."""
    cleaned = _strip_comments(content)
    cleaned = _strip_latex_environments(cleaned)
    cleaned = re.sub(r"\\(?:vspace|hspace|fontsize|selectfont)\*?\{[^}]*\}", " ", cleaned)
    cleaned = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", " ", cleaned)
    cleaned = re.sub(r"[{}\$^~]", " ", cleaned)
    return bool(cleaned.strip())


def convert_tex_to_docx(tex_path: Path, out_path: Path) -> None:
    """Convert a resume LaTeX file to a high-fidelity Word (.docx) document."""
    raw_tex = _read_tex(tex_path)
    tex = _expand_newcommands(_strip_comments(raw_tex))
    body = _extract_body(tex)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Standard ATS Style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(9.5)
    style.paragraph_format.line_spacing = 1.05
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # 1. Parse Candidate Header
    _parse_heading(doc, body)

    # 2. Parse Sections
    for section_match in re.finditer(r"\\section\{([^{}]+)\}(.*?)(?=\\section\{|\Z)", body, re.DOTALL):
        title = section_match.group(1).strip()
        content = section_match.group(2)
        if not _has_section_content(content):
            continue
        _add_section_heading(doc, title)
        _parse_section_content(doc, content, section_title=title)

    doc.save(out_path)
